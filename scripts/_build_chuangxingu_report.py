# -*- coding: utf-8 -*-
"""组装创维创新谷 10 月报告。优化前=2026-09-16，10月=2026-10-28。监测用客户项目 721。"""
import json
from datetime import datetime, date
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parent.parent
ENV = {}
for line in Path(r"e:\6.cursor\reports\report_gujinggong\.env").read_text(encoding="utf-8").splitlines():
    t = line.strip()
    if not t or t.startswith("#") or "=" not in t:
        continue
    k, v = t.split("=", 1)
    ENV[k.strip()] = v.strip().strip('"').strip("'")

PROJECT_ID = 721
PROJECT_NAME = "创新谷-优化词"
BEFORE = "2026-09-16"
OCTOBER = "2026-10-28"
DISPLAY = "创维创新谷"
BASE = ENV["GEO_API_BASE"].rstrip("/")
XLSX = Path(
    r"e:\02GEO\3.GEO数据监测系统\投放文章统计-1029\投放文章统计-1029"
    r"\创新谷-优化词(中位数)_投放文章统计_1029.xlsx"
)

DOMAIN_NAMES = {
    "mp.weixin.qq.com": "微信公众号",
    "post.smzdm.com": "什么值得买",
    "smzdm.com": "什么值得买",
    "hebei.ifeng.com": "凤凰网",
    "ishare.ifeng.com": "凤凰网",
    "ifeng.com": "凤凰网",
    "sz.ifeng.com": "凤凰网",
    "hainan.ifeng.com": "凤凰网",
    "sohu.com": "搜狐",
    "3g.k.sohu.com": "搜狐",
    "focus.cn": "搜狐焦点",
    "v.douyin.com": "抖音",
    "douyin.com": "抖音",
    "toutiao.com": "今日头条",
    "pad.zol.com.cn": "中关村在线",
    "news.zol.com.cn": "中关村在线",
    "xa.zol.com.cn": "中关村在线",
    "jd.zol.com.cn": "中关村在线",
    "price.zol.com.cn": "中关村在线",
    "xinwen.pconline.com.cn": "太平洋电脑网",
    "super.sina.cn": "新浪",
    "k.sina.com.cn": "新浪",
    "zhihu.com": "知乎",
    "view.inews.qq.com": "腾讯新闻",
    "news.qq.com": "腾讯新闻",
    "ithome.com": "IT之家",
    "baijiahao.baidu.com": "百家号",
    "news.cngold.org": "金投网",
    "weibo.com": "微博",
    "163.com": "网易",
    "itbear.com.cn": "ITBEAR",
    "software.it168.com": "IT168",
    "bjnews.com.cn": "新京报",
    "news.yesky.com": "天极网",
    "jianghu.taobao.com": "淘宝江湖",
    "dutenews.com": "读特",
    "cet.com.cn": "中国经济时报",
    "cet.com.cn": "中国经济时报",
    "chinaz.com": "站长之家",
    "tidenews.com.cn": "潮新闻",
    "cheari.com": "家电院",
    "digitalhome.cheaa.com": "中国家电网",
    "douban.com": "豆瓣",
    "example.com": "其他",
}

s = requests.Session()
login = s.post(
    f"{BASE}/login",
    json={"username": ENV["GEO_USER"], "password": ENV["GEO_PASS"]},
    timeout=60,
).json()
if not login.get("ok"):
    raise SystemExit("login failed")


def get(path, **params):
    r = s.get(f"{BASE}{path}", params=params, timeout=120)
    j = r.json()
    if not j.get("ok"):
        raise SystemExit(f"{path} {j.get('error')}")
    return j["data"]


def num(v):
    if v is None or v == "":
        return None
    return float(v)


def clean_num(v):
    n = num(v)
    if n is None:
        return None
    if float(n).is_integer():
        return int(n)
    return round(float(n), 1)


def page_all(path, **params):
    rows, page = [], 1
    total = None
    while page <= 30:
        data = get(path, page=page, page_size=100, **params)
        batch = data.get("list") or []
        rows.extend(batch)
        total = data.get("total", total)
        total_pages = data.get("total_pages") or 1
        if page >= total_pages or not batch:
            return rows, total if total is not None else len(rows)
        page += 1
    return rows, len(rows)


projects = get("/api/projects")
project = next((p for p in projects if p.get("id") == PROJECT_ID), None)
if not project or project.get("project_name") != PROJECT_NAME:
    raise SystemExit(f"项目不匹配: {project}")
if "中位数" in (project.get("project_name") or ""):
    raise SystemExit("拒绝使用中位数项目")

dates = []
try:
    dates = sorted((get(f"/api/projects/{PROJECT_ID}/data-dates").get("dates") or []))
except SystemExit:
    dates = []
print("project", project["id"], project["project_name"], project.get("target_product"))
print("has_before", BEFORE in dates, "has_october", OCTOBER in dates)
print("date_span", dates[0] if dates else None, dates[-1] if dates else None)

platforms = get("/api/platforms", project_id=PROJECT_ID)
plat_map = {p["id"]: p for p in platforms}
print("platforms", [(p["id"], p["name"]) for p in platforms])


def with_self(rows, n):
    head = list(rows[:n])
    me = next((b for b in rows if b.get("is_target") or b.get("is_self")), None)
    if me and me not in head:
        head.append(me)
    return head


def snapshot(day):
    stats = get("/api/conversations/stats", project_id=PROJECT_ID, start_date=day, end_date=day)
    influence = get("/api/competitors/influence", project_id=PROJECT_ID, start_date=day, end_date=day)
    entries, total = page_all(
        "/api/entries",
        project_id=PROJECT_ID,
        start_date=day,
        end_date=day,
        sort_by="mention_rate",
        sort_order="desc",
    )
    mention = page_all(
        "/api/competitors/mention-rate",
        project_id=PROJECT_ID,
        start_date=day,
        end_date=day,
    )[0]
    top1 = get(
        "/api/competitors/top-mention-rate",
        project_id=PROJECT_ID,
        start_date=day,
        end_date=day,
        top_type="top1",
        page_size=100,
    )
    top3 = get(
        "/api/competitors/top-mention-rate",
        project_id=PROJECT_ID,
        start_date=day,
        end_date=day,
        top_type="top3",
        page_size=100,
    )
    me = next((b for b in (influence.get("list") or []) if b.get("is_target")), {})
    top1_list = top1.get("list") or []
    top3_list = top3.get("list") or []
    self_top1 = next((b for b in top1_list if b.get("is_self") or b.get("is_target")), None)
    self_top3 = next((b for b in top3_list if b.get("is_self") or b.get("is_target")), None)
    self_mention = next((b for b in mention if b.get("is_target") or b.get("is_self")), None)

    def pname(p):
        return plat_map.get(p.get("platform_id"), {}).get("name") or p.get("platform_name") or f"平台{p.get('platform_id')}"

    def brand_name(b):
        if b.get("is_target") or b.get("is_self"):
            return DISPLAY
        return b.get("display_name") or b.get("product_name") or b.get("brand_name") or "-"

    def pack_rank(rows, rate_key):
        out = []
        for i, b in enumerate(rows):
            rate = b.get(rate_key)
            if rate is None:
                rate = b.get("selected_top_mention_rate")
            out.append({
                "name": brand_name(b),
                "value": None if rate is None else clean_num(rate),
                "is_target": bool(b.get("is_target") or b.get("is_self")),
                "rank": b.get("rank") or i + 1,
            })
        return out

    covered = [e for e in entries if (num(e.get("mention_rate")) or 0) > 0]
    return {
        "day": day,
        "mention_rate": clean_num(stats.get("brand_mention_rate")),
        "stats_top1": clean_num(stats.get("top1_mention_rate")),
        "stats_top3": clean_num(stats.get("top3_mention_rate")),
        "top1_mention_rate": clean_num(self_top1.get("selected_top_mention_rate")) if self_top1 else 0,
        "top3_mention_rate": clean_num(self_top3.get("selected_top_mention_rate")) if self_top3 else 0,
        "top1_rank": (self_top1 or {}).get("rank"),
        "top3_rank": (self_top3 or {}).get("rank"),
        "mention_rank": (self_mention or {}).get("rank"),
        "influence_score": clean_num(me.get("influence_score")),
        "influence_rank": me.get("rank"),
        "conversations": influence.get("total_conversations"),
        "entries_total": total or len(entries),
        "covered_entries": len(covered),
        "platforms": [
            {"name": pname(p), "mention_rate": clean_num(p.get("brand_mention_rate"))}
            for p in (stats.get("platform_stats") or [])
        ],
        "mention_board": pack_rank(with_self(mention, 8), "mention_rate"),
        "top1_board": pack_rank(with_self(top1_list, 8), "selected_top_mention_rate"),
        "top3_board": pack_rank(with_self(top3_list, 8), "selected_top_mention_rate"),
        "influence_board": pack_rank(with_self(influence.get("list") or [], 8), "influence_score"),
        "entries": [
            {
                "name": e.get("entry_name"),
                "mention_rate": clean_num(e.get("mention_rate")),
                "top1_mention_rate": clean_num(e.get("top1_mention_rate")),
            }
            for e in entries
        ],
    }


def fmt_rate(v):
    if v is None or v == "":
        return "--"
    return f"{float(v):.1f}%"


def fmt_date(v):
    if not v:
        return "--"
    text = str(v)[:10]
    y, m, d = text.split("-")
    return f"{int(y)}/{int(m)}/{int(d)}"


before = snapshot(BEFORE)
october = snapshot(OCTOBER)

platform_entries = {}
for pl in platforms:
    rows, _ = page_all(
        "/api/entries",
        project_id=PROJECT_ID,
        start_date=OCTOBER,
        end_date=OCTOBER,
        sort_by="mention_rate",
        sort_order="desc",
        platform_ids=pl["id"],
    )
    platform_entries[str(pl["id"])] = [
        {
            "entry_name": e.get("entry_name"),
            "mention_rate": fmt_rate(e.get("mention_rate")),
            "top1_mention_rate": fmt_rate(e.get("top1_mention_rate")),
            "last_conversation_time": fmt_date(e.get("last_conversation_time")),
            "last_screenshot_url": e.get("last_screenshot_url") or None,
        }
        for e in rows
    ]
    print(f"platform {pl['id']} {pl['name']}: {len(rows)}")

before_map = {e["name"]: e for e in before["entries"]}
entries = []
for e in october["entries"]:
    prev = before_map.get(e["name"]) or {}
    entries.append({
        "name": e["name"],
        "mention_rate": e["mention_rate"],
        "top1_mention_rate": e["top1_mention_rate"],
        "before_mention_rate": prev.get("mention_rate"),
        "before_top1_mention_rate": prev.get("top1_mention_rate"),
    })


def top_four_plus_self(board, value_fmt):
    ordered = sorted(board, key=lambda b: b["rank"] or 999)
    me = next((b for b in ordered if b["is_target"]), None)
    if me and (me["rank"] or 999) <= 5:
        rows = ordered[:5]
        if me not in rows:
            rows = ordered[:4] + [me]
    else:
        leaders = [b for b in ordered if not b["is_target"]][:4]
        rows = leaders + ([me] if me else [])
    return [
        {
            "name": b["name"],
            "value": value_fmt(b),
            "isTarget": b["is_target"],
            "rank": b["rank"],
        }
        for b in rows
    ]


def pct(v):
    return "-" if v is None else f"{v}%"


snap_keys = (
    "mention_rate", "top1_mention_rate", "top3_mention_rate", "influence_score",
    "influence_rank", "conversations", "covered_entries", "platforms",
)

report = {
    "meta": {
        "project_id": PROJECT_ID,
        "project_name": PROJECT_NAME,
        "product": DISPLAY,
        "brand": DISPLAY,
        "label": "2026年10月",
        "before_date": BEFORE,
        "september_date": OCTOBER,
        "october_date": OCTOBER,
        "before_label": "优化前",
        "september_label": "10月",
    },
    "scope": {
        "platforms": len(platforms),
        "platform_names": [p["name"] for p in platforms],
        "entries": october["entries_total"],
        "conversations": october["conversations"],
    },
    "before": {k: before[k] for k in snap_keys},
    "september": {k: october[k] for k in snap_keys},
    "compare": {
        "mention_rate": top_four_plus_self(october["mention_board"], lambda b: pct(b["value"])),
        "top1": top_four_plus_self(october["top1_board"], lambda b: pct(b["value"])),
        "influence": top_four_plus_self(october["influence_board"], lambda b: f"NO. {b['rank']}"),
    },
    "entries": entries,
    "platform_entries": platform_entries,
    "platforms": [{"id": p["id"], "name": p["name"]} for p in platforms],
}

(ROOT / "src" / "data" / "chuangxinguOctoberReport.json").write_text(
    json.dumps(report, ensure_ascii=False, indent=2),
    encoding="utf-8",
)


def as_date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value or "").strip()[:10]


wb = openpyxl.load_workbook(XLSX, data_only=True)
overview = wb["数据概览"]
detail = wb["文章引用明细"]
channels = wb["平台与渠道分析"]


def cell(row, col):
    return overview.cell(row, col).value


rows = []
for record in detail.iter_rows(min_row=2, values_only=True):
    if record[0] in (None, ""):
        continue
    try:
        int(record[0])
    except (TypeError, ValueError):
        continue
    rows.append(record)

delivery_articles = int(cell(14, 2))
cited_articles = int(cell(14, 5))
delivery_citations = int(cell(10, 2))
if len(rows) != delivery_articles:
    raise SystemExit(f"明细 {len(rows)} 行，概览投放文章总数 {delivery_articles}")
cited_rows = [r for r in rows if r[8] == "是"]
if len(cited_rows) != cited_articles:
    raise SystemExit(f"被引 {len(cited_rows)} 行，概览被引用文章数 {cited_articles}")


def total_col(idx):
    return sum(int(r[idx] or 0) for r in rows)


platform_totals = {
    "total": total_col(5),
    "deepseek": total_col(6),
    "doubao": total_col(7),
}
if platform_totals["total"] != delivery_citations:
    raise SystemExit(f"明细引用合计 {platform_totals['total']}，概览 {delivery_citations}")

unknown_domains = sorted({r[3] for r in rows if r[3] not in DOMAIN_NAMES})

grouped = {}
for r in rows:
    name = DOMAIN_NAMES.get(r[3], r[3])
    acc = grouped.setdefault(name, {"channel": name, "articles": 0, "cited": 0, "citations": 0})
    acc["articles"] += 1
    if r[8] == "是":
        acc["cited"] += 1
    acc["citations"] += int(r[5] or 0)
domain_stats = sorted(grouped.values(), key=lambda a: (a["citations"], a["articles"]), reverse=True)

top_channels = []
for record in channels.iter_rows(min_row=2, values_only=True):
    if not record[0] or record[0] == "发布平台":
        continue
    top_channels.append({
        "channel": record[0],
        "articles": int(record[1] or 0),
        "cited": int(record[2] or 0),
        "citedRate": float(record[3] or 0),
        "citations": int(record[4] or 0),
    })
top_channels.sort(key=lambda a: (a["citations"], a["articles"]), reverse=True)

delivery = {
    "source": XLSX.name,
    "generated_at": "2026-10-29",
    "overview": {
        "delivery_articles": delivery_articles,
        "cited_articles": cited_articles,
        "cited_rate": round(cited_articles / delivery_articles * 100, 1),
        "unique_titles": len({r[1] for r in rows}),
        "unique_cited_titles": len({r[1] for r in cited_rows}),
        "delivery_citations": delivery_citations,
        "total_conversations": int(cell(6, 2)),
        "total_citation_records": int(cell(6, 5)),
    },
    "platform_totals": platform_totals,
    "top10": [
        {
            "rank": int(r[0]),
            "title": r[1],
            "platform": r[2],
            "domain": r[3],
            "date": as_date(r[4]),
            "total": int(r[5] or 0),
            "deepseek": int(r[6] or 0),
            "doubao": int(r[7] or 0),
            "yuanbao": int(r[8] or 0),
            "isCited": r[8],
        }
        for r in rows[:10]
    ],
    "top_channels": top_channels,
    "domain_stats": domain_stats,
}
(ROOT / "src" / "data" / "chuangxinguDelivery.json").write_text(
    json.dumps(delivery, ensure_ascii=False, indent=2),
    encoding="utf-8",
)

zeros = [e for e in entries if not (e["mention_rate"] or 0)]
gains = sorted(entries, key=lambda e: (e["mention_rate"] or 0) - (e["before_mention_rate"] or 0), reverse=True)
top1_zero = [e for e in entries if (e["mention_rate"] or 0) >= 50 and not (e["top1_mention_rate"] or 0)]
new_cover = [e for e in entries if not (e["before_mention_rate"] or 0) and (e["mention_rate"] or 0) > 0]


def brief(snap):
    return {k: snap[k] for k in snap if k not in ("entries", "mention_board", "top1_board", "top3_board", "influence_board")}


summary = {
    "before": brief(before),
    "october": brief(october),
    "boards": {
        "mention": october["mention_board"],
        "top1": october["top1_board"],
        "top3": october["top3_board"],
        "influence": october["influence_board"],
        "before_mention": before["mention_board"][:6],
        "before_top1": before["top1_board"][:6],
    },
    "compare": report["compare"],
    "zero_count": len(zeros),
    "zero_names": [e["name"] for e in zeros],
    "new_cover": [
        {"name": e["name"], "now": e["mention_rate"], "top1": e["top1_mention_rate"]}
        for e in sorted(new_cover, key=lambda e: e["mention_rate"] or 0, reverse=True)[:12]
    ],
    "top_gains": [
        {"name": e["name"], "before": e["before_mention_rate"], "now": e["mention_rate"], "top1": e["top1_mention_rate"]}
        for e in gains[:10]
    ],
    "top1_zero_high_mention": [
        {"name": e["name"], "now": e["mention_rate"]}
        for e in sorted(top1_zero, key=lambda e: e["mention_rate"] or 0, reverse=True)[:10]
    ],
    "lowest": [
        {"name": e["name"], "now": e["mention_rate"], "top1": e["top1_mention_rate"], "before": e["before_mention_rate"]}
        for e in sorted(entries, key=lambda e: (e["mention_rate"] or 0, e["top1_mention_rate"] or 0))[:8]
    ],
    "platform_counts": {k: len(v) for k, v in platform_entries.items()},
    "delivery": {
        "articles": delivery_articles,
        "cited": cited_articles,
        "rate": delivery["overview"]["cited_rate"],
        "titles": delivery["overview"]["unique_titles"],
        "platforms": platform_totals,
        "channels": domain_stats[:8],
        "unknown_domains": unknown_domains,
        "top1_title": delivery["top10"][0]["title"] if delivery["top10"] else None,
    },
}
(ROOT / "scripts" / "_chuangxingu_summary.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2),
    encoding="utf-8",
)
print("entries", len(entries), "zeros", len(zeros), "articles", delivery_articles)
