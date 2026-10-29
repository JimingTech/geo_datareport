# -*- coding: utf-8 -*-
"""投放明细：创新谷-优化词(中位数) 1029 表。监测平台只有 DeepSeek / 豆包。"""
import json
from datetime import datetime, date
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parent.parent
ENV = {}
for line in Path(r"e:\6.cursor\reports\report_gujinggong\.env").read_text(encoding="utf-8").splitlines():
    t = line.strip()
    if not t or t.startswith("#") or "=" not in t:
        continue
    k, v = t.split("=", 1)
    ENV[k.strip()] = v.strip().strip('"').strip("'")

XLSX = Path(
    r"e:\02GEO\3.GEO数据监测系统\投放文章统计-1029\投放文章统计-1029"
    r"\创新谷-优化词(中位数)_投放文章统计_1029.xlsx"
)

DOMAIN_NAMES = {
    "mp.weixin.qq.com": "微信公众号",
    "post.smzdm.com": "什么值得买",
    "smzdm.com": "什么值得买",
    "hebei.ifeng.com": "凤凰网",
    "ishare.ifeng.com": "凤凰网",
    "ifeng.com": "凤凰网",
    "sz.ifeng.com": "凤凰网",
    "hainan.ifeng.com": "凤凰网",
    "sohu.com": "搜狐",
    "3g.k.sohu.com": "搜狐",
    "focus.cn": "搜狐焦点",
    "v.douyin.com": "抖音",
    "douyin.com": "抖音",
    "www.douyin.com": "抖音",
    "toutiao.com": "今日头条",
    "www.toutiao.com": "今日头条",
    "zhihu.com": "知乎",
    "www.zhihu.com": "知乎",
    "zhuanlan.zhihu.com": "知乎",
    "baijiahao.baidu.com": "百家号",
    "163.com": "网易",
    "www.163.com": "网易",
    "c.m.163.com": "网易",
    "news.163.com": "网易",
    "k.sina.com.cn": "新浪",
    "super.sina.cn": "新浪",
    "finance.sina.com.cn": "新浪",
    "news.sina.com.cn": "新浪",
    "jianghu.taobao.com": "淘宝江湖",
    "dutenews.com": "读特",
    "cet.com.cn": "中国经济时报",
    "www.cet.com.cn": "中国经济时报",
    "view.inews.qq.com": "腾讯新闻",
    "news.qq.com": "腾讯新闻",
    "weibo.com": "微博",
    "example.com": "其他",
}


def as_date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value or "").strip()[:10]


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    overview = wb["数据概览"]
    detail = wb["文章引用明细"]

    def cell(row, col):
        return overview.cell(row, col).value

    rows = []
    for record in detail.iter_rows(min_row=2, values_only=True):
        if record[0] in (None, ""):
            continue
        try:
            int(record[0])
        except (TypeError, ValueError):
            continue
        rows.append(record)

    delivery_articles = int(cell(14, 2))
    cited_articles = int(cell(14, 5))
    delivery_citations = int(cell(10, 2))
    if len(rows) != delivery_articles:
        raise SystemExit(f"明细 {len(rows)} 行，概览投放文章总数 {delivery_articles}")
    cited_rows = [r for r in rows if r[8] == "是"]
    if len(cited_rows) != cited_articles:
        raise SystemExit(f"被引 {len(cited_rows)} 行，概览被引用文章数 {cited_articles}")

    def total(idx):
        return sum(int(r[idx] or 0) for r in rows)

    platform_totals = {
        "total": total(5),
        "deepseek": total(6),
        "doubao": total(7),
    }
    if platform_totals["total"] != delivery_citations:
        raise SystemExit(f"明细引用合计 {platform_totals['total']}，概览 {delivery_citations}")

    unknown = sorted({r[3] for r in rows if r[3] not in DOMAIN_NAMES})
    grouped = {}
    for r in rows:
        name = DOMAIN_NAMES.get(r[3], r[3])
        acc = grouped.setdefault(name, {"channel": name, "articles": 0, "cited": 0, "citations": 0})
        acc["articles"] += 1
        if r[8] == "是":
            acc["cited"] += 1
        acc["citations"] += int(r[5] or 0)
    domain_stats = sorted(grouped.values(), key=lambda a: (a["citations"], a["articles"]), reverse=True)

    payload = {
        "source": XLSX.name,
        "generated_at": "2026-10-29",
        "overview": {
            "delivery_articles": delivery_articles,
            "cited_articles": cited_articles,
            "cited_rate": round(cited_articles / delivery_articles * 100, 1),
            "unique_titles": len({r[1] for r in rows}),
            "unique_cited_titles": len({r[1] for r in cited_rows}),
            "delivery_citations": delivery_citations,
            "total_conversations": int(cell(6, 2)),
            "total_citation_records": int(cell(6, 5)),
        },
        "platform_totals": platform_totals,
        "top10": [
            {
                "rank": int(r[0]),
                "title": r[1],
                "platform": r[2],
                "domain": r[3],
                "date": as_date(r[4]),
                "total": int(r[5] or 0),
                "deepseek": int(r[6] or 0),
                "doubao": int(r[7] or 0),
                "isCited": r[8],
            }
            for r in rows[:10]
        ],
        "domain_stats": domain_stats,
    }
    out = ROOT / "src" / "data" / "chuangxinguDelivery.json"
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    by_articles = sorted(domain_stats, key=lambda a: (a["articles"], a["citations"]), reverse=True)
    by_cited = sorted([a for a in domain_stats if a["cited"] > 0], key=lambda a: (a["cited"], a["citations"]), reverse=True)
    print(json.dumps({
        "articles": delivery_articles,
        "cited": cited_articles,
        "rate": payload["overview"]["cited_rate"],
        "titles": payload["overview"]["unique_titles"],
        "platforms": platform_totals,
        "channels": [f"{a['channel']} {a['articles']}篇/{a['cited']}引" for a in by_articles[:10]],
        "cited_channels": [f"{a['channel']} {a['cited']}" for a in by_cited[:6]],
        "unknown": unknown,
        "top3": [t["title"] for t in payload["top10"][:3]],
    }, ensure_ascii=False, indent=2))


def ranks():
    s = requests.Session()
    base = ENV["GEO_API_BASE"].rstrip("/")
    login = s.post(f"{base}/login", json={"username": ENV["GEO_USER"], "password": ENV["GEO_PASS"]}, timeout=60).json()
    if not login.get("ok"):
        raise SystemExit("login failed")

    def get(path, **params):
        j = s.get(f"{base}{path}", params=params, timeout=90).json()
        if not j.get("ok"):
            raise SystemExit(j.get("error"))
        return j["data"]

    for day in ("2026-09-16", "2026-10-28"):
        mention = get("/api/competitors/mention-rate", project_id=721, start_date=day, end_date=day, page=1, page_size=100)
        top1 = get("/api/competitors/top-mention-rate", project_id=721, start_date=day, end_date=day, top_type="top1", page_size=20)
        top3 = get("/api/competitors/top-mention-rate", project_id=721, start_date=day, end_date=day, top_type="top3", page_size=20)
        stats = get("/api/conversations/stats", project_id=721, start_date=day, end_date=day)
        me_m = next(b for b in (mention.get("list") or mention) if b.get("is_target") or b.get("is_self"))
        me_1 = next(b for b in top1["list"] if b.get("is_self") or b.get("is_target"))
        me_3 = next(b for b in top3["list"] if b.get("is_self") or b.get("is_target"))
        print(day, {
            "stats_mention": stats.get("brand_mention_rate"),
            "stats_top1": stats.get("top1_mention_rate"),
            "stats_top3": stats.get("top3_mention_rate"),
            "mention_rank": me_m.get("rank"),
            "mention_rate": me_m.get("mention_rate"),
            "top1_rank": me_1.get("rank"),
            "top1": me_1.get("selected_top_mention_rate"),
            "top3_rank": me_3.get("rank"),
            "top3": me_3.get("selected_top_mention_rate"),
        })


if __name__ == "__main__":
    main()
    ranks()
